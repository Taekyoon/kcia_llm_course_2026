"""PPTX 슬라이드의 본문과 발표자 노트를 마크다운으로 추출한다.

슬라이드마다 "쉬운 HPC 활용교육 HPC E+Z!" 보일러플레이트가 반복되므로 걸러낸다.
발표자 노트는 강의 흐름 파악에 필요하므로 반드시 포함한다. (CLAUDE.md §6 참조)
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

_BOILERPLATE = re.compile(r"쉬운\s*HPC\s*활용교육|^HPC\s*$|^E\s*$|^\+\s*$|^Z\s*$|^!\s*$")
_WS = re.compile(r"[ \t ]+")


def _clean(text: str) -> str:
    return _WS.sub(" ", text).strip()


def _shape_lines(shape) -> list[str]:
    if not shape.has_text_frame:
        return []
    lines = []
    for para in shape.text_frame.paragraphs:
        # <a:br/> 는 run 이 아니라서 그냥 join 하면 앞뒤 줄이 붙어 버린다.
        # (실제 슬라이드는 멀쩡한데 추출본만 깨져 보여 오판한 적이 있다 — 2026-08-27)
        t = _clean(para.text.replace(chr(11), chr(10))).replace(chr(10), chr(10) + "    ")
        if t and not _BOILERPLATE.search(t):
            lines.append(("    " * para.level) + t)
    return lines


def _notes(slide) -> list[str]:
    if not slide.has_notes_slide:
        return []
    tf = slide.notes_slide.notes_text_frame
    if tf is None:
        return []
    return [t for t in (_clean(p.text) for p in tf.paragraphs) if t]


def _sort_key(shape):
    """읽는 순서(위→아래, 좌→우)에 가깝게 도형을 정렬한다."""
    top = shape.top if shape.top is not None else Emu(0)
    left = shape.left if shape.left is not None else Emu(0)
    return (int(top), int(left))


def to_markdown(pptx_path: Path) -> str:
    prs = Presentation(str(pptx_path))
    out: list[str] = [
        f"# {pptx_path.stem}",
        "",
        f"> 원본: `{pptx_path.name}`  ·  슬라이드 {len(prs.slides)}장",
        "> 추출: `tools/extract_pptx.py`",
        "",
    ]

    for idx, slide in enumerate(prs.slides, start=1):
        body: list[str] = []
        for shape in sorted(slide.shapes, key=_sort_key):
            body.extend(_shape_lines(shape))
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [_clean(c.text) for c in row.cells]
                    body.append("| " + " | ".join(cells) + " |")

        out.append(f"## p{idx}")
        out.append("")
        out.extend(body if body else ["_(텍스트 없음 — 이미지 전용 슬라이드로 추정)_"])

        notes = _notes(slide)
        if notes:
            out.append("")
            out.append("**[발표자 노트]**")
            out.extend(f"> {n}" for n in notes)
        out.append("")

    return "\n".join(out) + "\n"


def main() -> None:
    ap = argparse.ArgumentParser(description="PPTX 본문·발표자 노트 추출")
    ap.add_argument("pptx", type=Path)
    ap.add_argument("-o", "--out", type=Path)
    args = ap.parse_args()

    md = to_markdown(args.pptx)
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(md, encoding="utf-8")
        print(f"wrote {args.out} ({len(md)} chars)")
    else:
        print(md)


if __name__ == "__main__":
    main()
