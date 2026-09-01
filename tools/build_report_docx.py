"""발주처 제출용 보고서 마크다운 → .docx 변환.

    uv run --with python-docx --project tools python tools/build_report_docx.py

**마크다운이 정본이다.** 내용 수정은 `review/10_사전셋팅내용.md` 에서 하고 이 스크립트를
다시 돌린다. docx 를 직접 고치면 다음 실행에 덮어써진다 (노트북·슬라이드와 같은 규율).

왜 python-docx 인가: 이 로컬 환경에는 node 도 LibreOffice 도 없다(CLAUDE.md §3-A).
docx 스킬의 표준 경로(docx npm)를 못 쓰므로 `uv run --with` 로 격리 설치해 쓴다.

한글(HWP)에서 .docx 를 열 수 있으므로, 사용자는 열어서 .hwp 로 저장하면 된다.
(.hwp 직접 쓰기는 COM 금지로 불가 — CLAUDE.md §4)

지원하는 마크다운은 이 문서가 쓰는 것만이다:
  # 제목 / > 부제 / ## 절 / 표 / 문단 / **굵게** / *기울임 각주* / ---
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "review" / "10_사전셋팅내용.md"
OUT = ROOT / "work" / "사전셋팅_내용.docx"
OUT_BODY = ROOT / "work" / "사전셋팅_본문.docx"   # 원본 HWP 제목 아래에 붙여넣을 본문만

# 발주처 원본 HWP 는 제목 상자(파란 그래픽)만 있는 백지다. 그 양식을 살리려면
# 제목을 다시 쓰지 말고 **본문만** 만들어 붙여넣어야 한다 (2026-09-01 사용자 지시).
BODY_ONLY = "--body-only" in sys.argv

FONT = "맑은 고딕"
BODY_PT = 9.5
TRACKING = -14         # 자간 -0.7pt
GRAY = RGBColor(0x59, 0x59, 0x59)



TITLE_BAR_TOP = "1B87C0"      # 원본 제목 상자 위쪽 진한 파랑
TITLE_BAR_BOTTOM = "7FD2E8"   # 아래쪽 하늘색


def title_bars(para, sz=36):
    """문단 위·아래에 굵은 색 막대를 넣는다 (원본 제목 상자 재현)."""
    pPr = para._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for side, color in (("top", TITLE_BAR_TOP), ("bottom", TITLE_BAR_BOTTOM)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "8")
        el.set(qn("w:color"), color)
        pbdr.append(el)
    pPr.append(pbdr)


def set_font(run, size=BODY_PT, bold=False, color=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # 한글은 eastAsia 폰트를 따로 지정해야 적용된다
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    # 자간 축소 (1/20 pt 단위) — 2쪽에 맞추기 위해 (2026-09-01 사용자 지시)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(TRACKING))
    run._element.rPr.append(sp)


def add_runs(para, text, size=BODY_PT, color=None):
    """**굵게** 만 처리한다."""
    for i, chunk in enumerate(text.split("**")):
        if not chunk:
            continue
        set_font(para.add_run(chunk), size=size, bold=(i % 2 == 1), color=color)


def is_block_start(s: str) -> bool:
    """새 블록의 시작인가. **굵게** 로 시작하는 본문 줄이 각주로 오인되지 않도록
    별표 한 개짜리만 각주로 본다 (2026-09-01 실측 버그)."""
    return s.startswith(("#", "|", ">", "---")) or (
        s.startswith("*") and not s.startswith("**"))


def shade(cell, hex_color="EDEDED"):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(el)


def add_table(doc, rows):
    header, *body = rows
    t = doc.add_table(rows=len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r, row in enumerate([header, *body]):
        for c, text in enumerate(row):
            cell = t.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, text, size=8)
            if r == 0:
                shade(cell)
                for run in p.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def parse_table(lines, i):
    """마크다운 표를 (행 목록, 다음 인덱스) 로."""
    rows = []
    while i < len(lines) and lines[i].startswith("|"):
        cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(set(c) <= set("-: ") for c in cells):   # 구분선은 버린다
            rows.append(cells)
        i += 1
    return rows, i


def build() -> None:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(3)
    pf.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)      # A4
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line or line.startswith("---"):
            i += 1
            continue

        if line.startswith("| "):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        if line.startswith("# ") and BODY_ONLY:
            i += 1
            continue

        if line.startswith("> ") and BODY_ONLY:
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.3
            set_font(p.add_run(line[2:]), size=14, bold=True)
            title_bars(p)

        elif line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(4)
            set_font(p.add_run(line[3:]), size=11.5, bold=True)

        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.2
            add_runs(p, line[2:], size=9, color=GRAY)

        elif line.startswith("*") and not line.startswith("**"):
            # 마지막 각주 — 여러 줄로 이어질 수 있다
            buf = [line]
            while not buf[-1].rstrip().endswith("*") and i + 1 < len(lines):
                i += 1
                buf.append(lines[i].rstrip())
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            run = p.add_run(" ".join(buf).strip("*"))
            set_font(run, size=9, color=GRAY)
            run.italic = True

        else:
            buf = [line]
            while i + 1 < len(lines) and lines[i + 1].strip() and \
                    not is_block_start(lines[i + 1]):
                i += 1
                buf.append(lines[i].rstrip())
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, " ".join(buf))

        i += 1

    out = OUT_BODY if BODY_ONLY else OUT
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    n_para = len(doc.paragraphs)
    print(f"생성: {out}")
    print(f"  문단 {n_para}개 · 표 {len(doc.tables)}개 · {out.stat().st_size:,} bytes")


if __name__ == "__main__":
    if not SRC.exists():
        raise SystemExit(f"[중단] 원고가 없습니다: {SRC}")
    build()
